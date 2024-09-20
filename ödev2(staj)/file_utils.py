import os
from werkzeug.utils import secure_filename
from docx import Document
import pandas as pd
from PyPDF2 import PdfReader
from openpyxl import load_workbook

ALLOWED_EXTENSIONS = {'txt', 'pdf', 'docx', 'xlsx'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def read_txt(file_path):
    with open(file_path, 'r', encoding='utf-8') as file:
        return file.read()

def read_docx(file_path):
    doc = Document(file_path)
    content = []
    yazar = 'Unknown'
    olusturma_tarihi = 'Unknown'
    sahibi = 'Unknown'

    core_props = doc.core_properties
    if core_props.author:
        yazar = core_props.author
    if core_props.created:
        olusturma_tarihi = core_props.created.strftime('%Y-%m-%d %H:%M:%S')
    if core_props.last_modified_by:
        sahibi = core_props.last_modified_by
    
    for para in doc.paragraphs:
        content.append(para.text)
    return '\n'.join(content), yazar, olusturma_tarihi, sahibi



def read_excel(file_path):
    workbook = load_workbook(file_path)
    props = workbook.properties

    yazar = props.creator or 'Unknown'
    olusturma_tarihi = props.created.strftime('%Y-%m-%d %H:%M:%S') if props.created else 'Unknown'
    sahibi = props.last_modified_by or 'Unknown'

    df = pd.read_excel(file_path)
    return df.to_string(index=False), yazar, olusturma_tarihi, sahibi


def read_pdf(file_path):
    content = []
    yazar = 'Unknown'
    olusturma_tarihi = 'Unknown'
    sahibi = 'Unknown'
    
    with open(file_path, 'rb') as file:
        reader = PdfReader(file)
        for page in reader.pages:
            content.append(page.extract_text())
        
        doc_info = reader.metadata
        if doc_info.get('/Author'):
            yazar = doc_info['/Author']
        if doc_info.get('/CreationDate'):
            creation_date = doc_info['/CreationDate']
            olusturma_tarihi = f"{creation_date[:4]}-{creation_date[4:6]}-{creation_date[6:8]} {creation_date[8:10]}:{creation_date[10:12]}:{creation_date[12:14]}"
        if doc_info.get('/Producer'):
            sahibi = doc_info['/Producer']
    
    return '\n'.join(content), yazar, olusturma_tarihi, sahibi
